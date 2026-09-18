import contextlib
import io
import json
from pathlib import Path
import tempfile
import unittest

import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import Workbook

from api.app import create_app
from src.ingestion.loader import carregar_dados, carregar_tabelas_arquivo
from src.ingestion.multi_loader import carregar_multiplas_tabelas
from src.ingestion.report_normalizer import normalizar_relatorio, numero, StructuralReviewRequired


HEADER = ['Pedido_ID', 'Data', 'Faturamento', 'Lucro']
ROWS = [[1, '2025-01-01', 100, 40], [2, '2025-01-02', 200, 50]]


class NormalizerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def excel(self, rows, sheets=None):
        p = self.root / 'entrada.xlsx'
        with pd.ExcelWriter(p) as writer:
            for name, data in (sheets or {'Dados': rows}).items():
                pd.DataFrame(data).to_excel(writer, sheet_name=name, index=False, header=False)
        return p

    def report(self, extra=None):
        rows = [HEADER, ['JANEIRO'], *ROWS, *(extra or []), ['TOTAL', None, 300, 90]]
        return normalizar_relatorio(pd.DataFrame(rows))

    def test_csv_convencional(self):
        p = self.root / 'normal.csv'
        pd.DataFrame(ROWS, columns=HEADER).to_csv(p, index=False)
        pd.testing.assert_frame_equal(carregar_dados(p), pd.read_csv(p))

    def test_excel_convencional(self):
        p = self.excel([HEADER, *ROWS])
        pd.testing.assert_frame_equal(carregar_dados(p), pd.read_excel(p))

    def test_sem_cabecalho_convencional(self):
        self.assertEqual(carregar_dados(self.excel(ROWS)).iloc[0, 0], 1)

    def test_titulo(self):
        r = self.report()
        p = self.excel([['RELATÓRIO OPERACIONAL'], HEADER, ['Janeiro'], *ROWS, ['TOTAL', None, 300]])
        df = carregar_dados(p)
        self.assertEqual(len(df), 2)
        self.assertEqual(df.attrs['ingestao']['normalizacao']['auditoria_linhas'][0]['tipo'], 'titulo')

    def test_linha_vazia(self):
        r = self.report([[None] * 4])
        self.assertEqual(r['metadados']['linhas_vazias'], 1)

    def test_meses_sem_posicao_fixa(self):
        for month in ['JANEIRO', 'MARÇO', 'DECEMBER']:
            r = self.report([[month], *ROWS])
            self.assertEqual(r['metadados']['linhas_periodo'], 2)

    def test_ano(self):
        self.assertEqual(self.report([[2031]])['metadados']['linhas_periodo'], 2)

    def test_trimestre(self):
        self.assertEqual(self.report([['1º TRIMESTRE']])['metadados']['linhas_periodo'], 2)

    def test_resumo(self):
        self.assertEqual(self.report([['Saldo', None, 300]])['metadados']['linhas_resumo'], 2)

    def test_meta(self):
        self.assertTrue(any(r['tipo'] == 'meta' for r in self.report([['Meta diária', None, 500]])['metadados']['resumos']))

    def test_subtotal(self):
        self.assertTrue(any(r['tipo'] == 'subtotal' for r in self.report([['Subtotal', None, 500]])['metadados']['resumos']))

    def test_total_nao_soma(self):
        self.assertEqual(self.report()['blocos'][0]['Faturamento'].sum(), 300)

    def test_percentual_resumo(self):
        r = self.report([['Percentual atingido', None, '80%']])
        self.assertEqual(r['metadados']['resumos'][0]['valores'][2], '80%')

    def test_moeda_separada(self):
        rows = [['Data', 'Cliente', None, 'Faturamento'], ['Janeiro'], ['2025-01-01', 'A', 'R$', '1.204,00'], ['2025-01-02', 'B', 'R$', '1204,00'], ['Total', None, 'R$', '2.408,00']]
        df = carregar_dados(self.excel(rows))
        self.assertEqual(df['Faturamento'].sum(), 2408)
        self.assertEqual(len(df.columns), 3)

    def test_negativos(self):
        self.assertEqual(numero('-R$ 35,35'), -35.35)
        self.assertEqual(numero('(1.200,00)'), -1200)

    def test_brasileiro(self):
        for v in ['R$ 1.204,00', '1.204,00', '1204,00']:
            self.assertEqual(numero(v), 1204)

    def test_americano(self):
        for v in ['$1,204.00', '1,204.00', '1204.00']:
            self.assertEqual(numero(v), 1204)

    def test_moedas_diversas(self):
        for symbol in ['£', '€', '¥', '₹']:
            self.assertEqual(numero(symbol + '1204.00'), 1204)

    def test_milhar_ambiguo(self):
        self.assertIsNone(numero('1.204'))
        self.assertIsNone(numero('1,204'))

    def test_blocos_repetidos(self):
        r = self.report([['Fevereiro'], *ROWS])
        self.assertEqual(r['metadados']['registros_extraidos'], 4)
        self.assertEqual(r['metadados']['quantidade_blocos'], 1)

    def test_cabecalho_diferente_separa(self):
        r = self.report([['Produto', 'Estoque', 'Custo', 'Quantidade'], ['A', 50, 20, 1], ['B', 30, 10, 2]])
        self.assertEqual(len(r['blocos']), 2)

    def test_varias_abas_compativeis(self):
        p = self.excel([], {'A': [HEADER, *ROWS], 'B': [HEADER, *ROWS]})
        df = carregar_dados(p)
        self.assertEqual(len(df), 4)
        self.assertEqual(len(df.attrs['ingestao']['continuacoes']), 2)

    def test_abas_incompativeis(self):
        p = self.excel([], {'Vendas': [HEADER, *ROWS], 'Produtos': [['Produto', 'Estoque'], ['A', 2], ['B', 3]]})
        self.assertEqual(len(carregar_tabelas_arquivo(p)), 2)
        with self.assertRaises(StructuralReviewRequired):
            carregar_dados(p)

    def test_aba_vazia(self):
        p = self.excel([], {'A': [HEADER, *ROWS], 'Vazia': []})
        self.assertEqual(len(carregar_dados(p)), 2)

    def test_aba_resumos_preservada(self):
        p = self.excel([], {'A': [HEADER, *ROWS], 'Metas': [['Meta', 300], ['Total', 200]]})
        df = carregar_dados(p)
        self.assertEqual(len(df.attrs['ingestao']['diagnosticos_abas']['Metas']['resumos']), 2)

    def test_titulo_mesclado(self):
        p = self.root / 'merged.xlsx'
        book = Workbook(); sheet = book.active
        sheet.append(['RELATÓRIO']); sheet.merge_cells('A1:D1')
        for row in [HEADER, ['Janeiro'], *ROWS, ['Total', None, 300]]: sheet.append(row)
        book.save(p); book.close()
        info = carregar_dados(p).attrs['ingestao']['normalizacao']
        self.assertEqual(info['evidencias_excel']['celulas_mescladas'], ['A1:D1'])

    def test_desconhecida_preservada(self):
        r = self.report([['texto sem classificação', None, None, None]])
        self.assertTrue(r['metadados']['revisao_necessaria'])
        self.assertEqual(r['metadados']['linhas_desconhecidas'], 1)
        self.assertTrue(any(row.get('valores', [None])[0] == 'texto sem classificação' for row in r['metadados']['auditoria_linhas']))

    def test_muitos_vazios_preserva_coluna(self):
        rows = [HEADER + ['Extra']] + [row + ['raro' if i == 0 else None] for i, row in enumerate(ROWS * 50)]
        self.assertIn('Extra', carregar_dados(self.excel(rows)))

    def test_comentario(self):
        self.assertFalse(self.report([['Observação: valores conferidos']])['metadados']['revisao_necessaria'])

    def test_cabecalho_repetido(self):
        r = self.report([HEADER, *ROWS])
        self.assertEqual(len(r['blocos']), 1)
        self.assertEqual(len(r['blocos'][0]), 4)

    def test_data_registro_prioritaria(self):
        r = self.report()
        row = next(r for r in r['metadados']['auditoria_linhas'] if r['tipo'] == 'transacao')
        self.assertEqual(row['datas_registro'], ['2025-01-01'])
        self.assertEqual(row['periodo_contexto'], 'JANEIRO')

    def test_pequena_variacao(self):
        r = self.report([[3, '2025-01-03', 100, None]])
        self.assertEqual(r['metadados']['registros_extraidos'], 3)

    def test_invalido(self):
        p = self.root / 'invalid.xlsx'; p.write_bytes(b'not excel')
        with self.assertRaises(ValueError): carregar_dados(p)

    def test_baixa_confianca_interrompe(self):
        p = self.excel([['Janeiro'], ['ruído'], ['???']])
        with self.assertRaises(StructuralReviewRequired): carregar_dados(p)

    def test_formula_sem_cache(self):
        p = self.excel([HEADER, ['Janeiro'], *ROWS, [3, '2025-01-03', '=100+200', 10]])
        with self.assertRaises(StructuralReviewRequired) as result: carregar_dados(p)
        self.assertIn('formulas', json.dumps(result.exception.diagnostico))

    def test_formula_total_preservada(self):
        df = carregar_dados(self.excel([HEADER, ['Janeiro'], *ROWS, ['Total', None, '=SUM(C3:C4)']]))
        self.assertEqual(len(df), 2)
        self.assertIn('=SUM', json.dumps(df.attrs['ingestao']))

    def test_json_estrito(self):
        r = self.report()
        json.dumps(r['metadados'], allow_nan=False)

    def test_texto_com_moeda_nao_e_valor(self):
        rows = [HEADER + ['Descrição'], *[r + ['FALTAM R$ 458,00'] for r in ROWS]]
        df = carregar_dados(self.excel(rows))
        self.assertNotIn('normalizacao', df.attrs['ingestao'])
        self.assertEqual(df['Descrição'].tolist(), ['FALTAM R$ 458,00'] * 2)

    def test_total_com_data_e_ambiguo(self):
        self.assertTrue(self.report([['Total', '2025-01-31', 300, 90]])['metadados']['revisao_necessaria'])

    def test_posicoes_diferentes(self):
        for ordem in ([3, 0, 2, 1], [2, 1, 3, 0]):
            rows = [[HEADER[c] for c in ordem], ['Janeiro'], *[[row[c] for c in ordem] for row in ROWS]]
            df = carregar_dados(self.excel(rows))
            self.assertEqual(df['Faturamento'].sum(), 300)

    def test_cabecalho_parcial(self):
        rows = [['Pedido_ID', 'Data', None, 'Lucro'], ['Janeiro'], *ROWS]
        df = carregar_dados(self.excel(rows))
        self.assertIn('coluna_3', df)
        self.assertNotIn('Faturamento', df)

    def test_cabecalho_duas_linhas_mescladas(self):
        p = self.root / 'header.xlsx'
        book = Workbook(); sheet = book.active
        sheet.append(['Cadastro', None, 'Financeiro', None])
        sheet.merge_cells('A1:B1'); sheet.merge_cells('C1:D1')
        for row in [HEADER, *ROWS]: sheet.append(row)
        book.save(p); book.close()
        df = carregar_dados(p)
        self.assertEqual(df.columns.tolist(), ['Cadastro / Pedido_ID', 'Cadastro / Data', 'Financeiro / Faturamento', 'Financeiro / Lucro'])
        self.assertEqual(len(df), 2)

    def test_regioes_laterais_preservadas_para_revisao(self):
        rows = [['Cliente', 'Valor', None, None, None], ['A', 10, None, 'Produto', 'Estoque'], ['B', 20, None, 'X', 2], [None, None, None, 'Y', 3]]
        r = normalizar_relatorio(pd.DataFrame(rows))
        self.assertEqual(len(r['blocos']), 2)
        self.assertTrue(r['metadados']['revisao_necessaria'])
        self.assertEqual(len(r['metadados']['linhas_preservadas']), 4)

    def test_mudanca_incompativel_sem_header_interrompe(self):
        self.assertTrue(self.report([['Outro', 90, 'Armazém', 'Nova região'], ['Diferente', 20, 'Depósito', 'Outra região']])['metadados']['revisao_necessaria'])

    def test_xls_complexo(self):
        import xlwt
        p = self.root / 'legado.xls'; book = xlwt.Workbook(); sheet = book.add_sheet('Dados')
        for i, row in enumerate([HEADER, ['Janeiro'], *ROWS, ['Total', None, 300]]):
            for j, value in enumerate(row):
                if value is not None: sheet.write(i, j, value)
        book.save(str(p))
        self.assertEqual(carregar_dados(p)['Faturamento'].sum(), 300)

    def test_cabecalho_depois_da_amostra_inicial(self):
        p = self.excel([['Relatório operacional']] * 35 + [HEADER, *ROWS])
        df = carregar_dados(p)
        self.assertEqual(len(df), 2)
        self.assertEqual(df.attrs['ingestao']['linha_cabecalho'], 36)

    def test_csv_complexo_linhas_larguras_diferentes(self):
        p = self.root / 'operacional.csv'
        p.write_text('Pedido_ID;Data;Faturamento;Lucro\nJaneiro\n1;2025-01-01;100;40\n2;2025-01-02;200;50\nTotal;;300;90\n', encoding='utf-8')
        self.assertEqual(carregar_dados(p)['Faturamento'].astype(float).sum(), 300)

    def test_relatorio_sem_header_nao_inventa_semantica(self):
        df = carregar_dados(self.excel([['Janeiro'], *ROWS, ['Total', None, 300]]))
        self.assertEqual(df.columns.tolist(), ['coluna_1', 'coluna_2', 'coluna_3', 'coluna_4'])
        self.assertEqual(len(df), 2)

    def test_datas_nativas_auditoria_serializavel(self):
        from datetime import datetime
        r = self.report([[3, datetime(2025, 1, 3), 100, 10]])
        self.assertFalse(r['metadados']['revisao_necessaria'])
        json.dumps(r['metadados'], allow_nan=False)

    def test_formula_sem_header_em_coluna_numerica(self):
        p = self.excel([*ROWS, [3, '2025-01-03', '=1+2', 40]])
        with self.assertRaises(StructuralReviewRequired) as result:
            carregar_dados(p)
        self.assertIn('=1+2', json.dumps(result.exception.diagnostico))

    def test_multiplos_arquivos(self):
        self.excel([HEADER, ['Janeiro'], *ROWS])
        pd.DataFrame({'Produto': ['A'], 'Estoque': [20]}).to_csv(self.root / 'produtos.csv', index=False)
        self.assertEqual(len(carregar_multiplas_tabelas(self.root, estrito=True)), 2)

    def test_post_complexo(self):
        p = self.excel([HEADER, ['Janeiro'], *ROWS, ['Total', None, 300]])
        with TestClient(create_app(self.root / 'reports/resumo_executivo.json', self.root / 'uploads')) as client, contextlib.redirect_stdout(io.StringIO()):
            result = client.post('/api/analysis', files={'files': (p.name, p.read_bytes())})
        self.assertEqual(result.status_code, 200, result.text)
        summary = result.json()['summary']
        self.assertEqual(summary['kpis']['faturamento_total'], 300)
        self.assertEqual(summary['dados']['ingestao']['normalizacao']['registros_extraidos'], 2)

    def test_post_revisao_preserva_diagnostico(self):
        p = self.excel([HEADER, ['Janeiro'], *ROWS, ['Texto inexplicável']])
        with TestClient(create_app(self.root / 'reports/resumo_executivo.json', self.root / 'uploads')) as client:
            result = client.post('/api/analysis', files={'files': (p.name, p.read_bytes())})
        self.assertEqual(result.status_code, 422)
        self.assertIn('diagnostico_estrutural', result.json())
        self.assertTrue((self.root / 'reports/diagnostico_estrutural.json').exists())
        self.assertFalse((self.root / 'reports/resumo_executivo.json').exists())


if __name__ == '__main__':
    unittest.main()
