from pathlib import Path
from openpyxl import load_workbook
from collections import Counter
import json,sys
sys.stdout.reconfigure(encoding='utf-8')
d=json.loads(Path('reports/diagnostico_estrutural.json').read_text(encoding='utf-8'))
a=d['abas']['Planilha1']['auditoria_linhas']
print('DIAG_TYPES',dict(Counter(x['tipo'] for x in a)))
print('DIAG_HEAD',[(x['linha'],x['valores']) for x in a if x['tipo'].startswith('cabecalho')][:10])
p=Path(r'C:\Users\User\Downloads\2017 a 2020.xlsx')
w=load_workbook(p,data_only=False,read_only=True).active
years=Counter(); colb=Counter(); ids=[]; all_date_rows=0; id_without_date=0; dates_in_trans=Counter()
for r in range(1,w.max_row+1):
    idv=w.cell(r,1).value; dt=w.cell(r,2).value
    if hasattr(dt,'year'):
        all_date_rows+=1; years[dt.year]+=1; colb[dt.year]+=1
        if isinstance(idv,(int,float)):
            ids.append(idv); dates_in_trans[dt.year]+=1
    elif isinstance(idv,(int,float)):
        id_without_date+=1
print('DATE_ROWS',all_date_rows,'DATE_YEAR_COUNTS',dict(sorted(years.items())))
print('ID_WITH_DATE',len(ids),'ID_UNIQUE',len(set(ids)),'ID_DUP',len(ids)-len(set(ids)),'NUMERIC_ID_WITHOUT_DATE',id_without_date)
# row classification confidence and semantic issue info
tx=[x for x in a if x['tipo']=='transacao']
print('NORMALIZER_TX',len(tx),'TX_HAS_DATE_AUDIT',sum(bool(x.get('datas_registro')) for x in tx),'TX_IDS_DIRECT',sum(isinstance(w.cell(x['linha'],1).value,(int,float)) for x in tx))
print('TX_WITHOUT_DATE',[(x['linha'],x['valores'][:4]) for x in tx if not x.get('datas_registro')][:12])
print('TX_ID_NONNUM',[(x['linha'],x.get('datas_registro'),x['valores'][:4]) for x in tx if not isinstance(w.cell(x['linha'],1).value,(int,float))][:12])
