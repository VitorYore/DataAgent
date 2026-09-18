from pathlib import Path
from openpyxl import load_workbook
import json, re, unicodedata
from collections import Counter
p=Path(r'C:\Users\User\Downloads\2017 a 2020.xlsx')
wb=load_workbook(p, data_only=False, read_only=False)
terms=['PEDIDO','DATA','VALOR TOTAL','VALOR C/ DESCONTO','FORMA PAGAMENTO','CLIENTE','MARGEM BRUTA','MARGEM BRUTA %']
def norm(v):
    s=unicodedata.normalize('NFKD', str(v or '')).encode('ascii','ignore').decode().upper()
    return re.sub(r'\s+',' ',s).strip()
for ws in wb.worksheets:
    rows=list(ws.iter_rows(values_only=True))
    nrow=max((i+1 for i,row in enumerate(rows) if any(v is not None for v in row)),default=0)
    ncol=max((j+1 for row in rows for j,v in enumerate(row) if v is not None),default=0)
    hits=[]; formula=[]; merged=list(ws.merged_cells.ranges); empty=[]
    for i,row in enumerate(rows[:nrow]):
        vals=[v for v in row[:ncol] if v is not None]
        if not vals: empty.append(i+1)
        for j,v in enumerate(row[:ncol]):
            if isinstance(v,str) and norm(v) in terms: hits.append({'row':i+1,'col':j+1,'term':norm(v)})
            if isinstance(v,str) and v.startswith('='): formula.append(f'{ws.title}!{ws.cell(i+1,j+1).coordinate}')
    print('\nSHEET',json.dumps(ws.title,ensure_ascii=False),'shape',nrow,ncol,'merges',len(merged),'formulas',len(formula),'blank_rows',len(empty))
    print('MERGES', [str(x) for x in merged[:12]], 'FORMULAS',formula[:12])
    print('TERM_COUNTS',json.dumps(dict(Counter(x['term'] for x in hits)),ensure_ascii=False))
    print('HITS',json.dumps(hits[:100],ensure_ascii=False))
    for h in hits:
        if h['term']=='PEDIDO':
            for r in range(max(1,h['row']-2),min(nrow,h['row']+5)+1):
                vals=[ws.cell(r,c).value for c in range(1,ncol+1)]
                while vals and vals[-1] is None: vals.pop()
                print('ROW',r,json.dumps(vals,ensure_ascii=False,default=str))
            print('...')
