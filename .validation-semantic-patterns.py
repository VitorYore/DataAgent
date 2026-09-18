from pathlib import Path
from openpyxl import load_workbook
from collections import Counter
import json, re
p=Path(r'C:\Users\User\Downloads\2017 a 2020.xlsx')
wb=load_workbook(p, data_only=False, read_only=False)
ws=wb.active
wd=load_workbook(p, data_only=True, read_only=True).active
def kind(v):
 if v is None: return '_'
 if isinstance(v,str) and v.startswith('='): return 'F'
 if hasattr(v,'year') and hasattr(v,'month'): return 'D'
 if isinstance(v,(int,float)): return 'N'
 return 'T'
def show(a,b):
 for r in range(a,b+1):
  vals=[ws.cell(r,c).value for c in range(1,15)]
  if any(v is not None for v in vals):
   print(r, ''.join(kind(v) for v in vals), json.dumps(vals,ensure_ascii=False,default=str))
print('BEGINNING');show(1,14)
print('PRE HEADER');show(3940,3957)
print('POST HEADER');show(3956,3970)
# Gather rows with first cell integer and date somewhere; summarize signatures
for lo,hi,name in [(1,3956,'before'),(3957,12681,'after')]:
 c=Counter(); examples={}; candidates=[]
 for r in range(lo,hi+1):
  vals=[ws.cell(r,k).value for k in range(1,15)]
  if isinstance(vals[0],(int,float)) and any(kind(v)=='D' for v in vals):
   sig=''.join(kind(v) for v in vals)
   c[sig]+=1; examples.setdefault(sig,(r,vals)); candidates.append(r)
 print(name,'id_date_rows',len(candidates),'signature_top',c.most_common(12))
 for sig,n in c.most_common(5): print(' SIG',sig,n,'EXAMPLE',examples[sig][0],json.dumps(examples[sig][1],ensure_ascii=False,default=str))
 # Count ID candidates and duplicate, valid date
 ids=[]
 for r in candidates:
  val=ws.cell(r,1).value
  if val is not None: ids.append(val)
 print('unique',len(set(ids)),'duplicates',len(ids)-len(set(ids)),'id_minmax',min(ids) if ids else None,max(ids) if ids else None)
 print('period/head rows near interval')
 for r in range(lo,hi+1):
  vals=[ws.cell(r,k).value for k in range(1,15)]
  texts=[str(v).strip() for v in vals if isinstance(v,str) and not v.startswith('=')]
  if len(texts)<=2 and texts and any(x.upper() in {'JANEIRO','FEVEREIRO','MARÇO','MARCO','ABRIL','MAIO','JUNHO','JULHO','AGOSTO','SETEMBRO','OUTUBRO','NOVEMBRO','DEZEMBRO'} for x in texts):
   if r<60 or r>3950: print(' PERIOD',r,texts)
# formula cached status
form=0; cached=0; vals=[]
for row in ws.iter_rows():
 for cell in row:
  if isinstance(cell.value,str) and cell.value.startswith('='):
   form+=1; cv=wd[cell.coordinate].value
   if cv is not None:
    cached+=1
    if cell.column==8 and isinstance(cv,(int,float)): vals.append(float(cv))
print('FORMULAS',form,'CACHED',cached,'H_PERCENT_CACHED',len(vals),'first',vals[:5])
